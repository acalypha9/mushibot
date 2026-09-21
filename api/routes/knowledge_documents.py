import base64, datetime, os, re, tempfile, urllib.parse, uuid
from pathlib import Path
from typing import List, Optional
from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import Response
from sqlalchemy import text
from sqlalchemy.orm import Session
from auth import require_admin
from database import get_db
from models import KnowledgeCollection, KnowledgeDocument, User
from routes.knowledge_collections import get_setting, touch_collection
from schemas import DocumentResponse, ManualDocumentCreate

router = APIRouter()

def make_doc_response(d: KnowledgeDocument) -> DocumentResponse:
    meta = d.cmetadata if isinstance(d.cmetadata, dict) else {}
    c = d.created_at.isoformat() if d.created_at else ""
    u = d.updated_at.isoformat() if d.updated_at else c
    return DocumentResponse(
        id=str(d.id), collection_id=str(d.collection_id), title=d.title,
        document_type=d.document_type, source_type=d.source_type, file_size=d.file_size or 0,
        status=d.status, ingestion_status=d.ingestion_status, chunk_count=d.chunk_count,
        version=d.version, created_at=c, updated_at=u,
        convert_to_md=bool(meta.get("convert_to_md", False)), save_parser_output=bool(meta.get("save_parser_output", False)),
    )

@router.get("/documents", response_model=List[DocumentResponse])
def list_documents(collection_id: Optional[str] = None, current_user: User = Depends(require_admin), db: Session = Depends(get_db)):
    q = db.query(KnowledgeDocument)
    if collection_id:
        try: q = q.filter(KnowledgeDocument.collection_id == uuid.UUID(collection_id))
        except ValueError: pass
    return [make_doc_response(d) for d in q.order_by(KnowledgeDocument.created_at.desc()).all()]

@router.post("/documents", response_model=DocumentResponse, status_code=201)
def create_manual_document(data: ManualDocumentCreate, current_user: User = Depends(require_admin), db: Session = Depends(get_db)):
    try: col_uuid = uuid.UUID(data.collection_id)
    except ValueError: raise HTTPException(status_code=400, detail="Invalid collection ID")
    col = db.query(KnowledgeCollection).filter(KnowledgeCollection.id == col_uuid).first()
    if not col: raise HTTPException(status_code=404, detail="Collection not found")
    doc = KnowledgeDocument(
        id=uuid.uuid4(), collection_id=col.id, title=data.title, document_type=data.document_type,
        source_type="MANUAL_TEXT", content_text=data.content_text, markdown_text=data.content_text,
        file_size=len(data.content_text.encode("utf-8")), status="DRAFT", ingestion_status="PENDING", chunk_count=0,
    )
    db.add(doc)
    db.commit()
    db.refresh(doc)
    return make_doc_response(doc)

@router.post("/documents/upload", response_model=DocumentResponse, status_code=201)
def upload_document(
    collection_id: str = Form(...), title: str = Form(...), document_type: str = Form("POLICY"),
    convert_to_md: bool = Form(True), save_parser_output: bool = Form(False),
    chunk_size: Optional[int] = Form(None), chunk_overlap: Optional[int] = Form(None),
    file: UploadFile = File(...), current_user: User = Depends(require_admin), db: Session = Depends(get_db),
):
    try: col_uuid = uuid.UUID(collection_id)
    except ValueError: raise HTTPException(status_code=400, detail="Invalid collection ID")
    col = db.query(KnowledgeCollection).filter(KnowledgeCollection.id == col_uuid).first()
    if not col: raise HTTPException(status_code=404, detail="Collection not found")
    content_bytes = file.file.read()
    raw_text = None
    if file.filename.lower().endswith((".txt", ".md", ".markdown", ".rst", ".adoc")):
        try: raw_text = content_bytes.decode("utf-8")
        except UnicodeDecodeError: raw_text = content_bytes.decode("latin-1", errors="ignore")
    if convert_to_md and title:
        title = re.sub(r"\.(pdf|docx|doc|txt|rst|adoc)$", ".md", title, flags=re.IGNORECASE)
        if not title.lower().endswith(".md"): title += ".md"
    doc = KnowledgeDocument(
        id=uuid.uuid4(), collection_id=col.id, title=title, document_type=document_type,
        source_type="FILE_UPLOAD", content_text=raw_text, markdown_text=raw_text,
        file_size=len(content_bytes), status="DRAFT", ingestion_status="PENDING", chunk_count=0,
        cmetadata={
            "convert_to_md": convert_to_md, "save_parser_output": save_parser_output,
            "chunk_size": chunk_size or 512, "chunk_overlap": chunk_overlap or 50,
            "filename": file.filename, "file_b64": base64.b64encode(content_bytes).decode("ascii"),
        },
    )
    db.add(doc)
    db.commit()
    db.refresh(doc)
    return make_doc_response(doc)

def _parse_content_to_markdown(doc: KnowledgeDocument, db: Session) -> str:
    from knowledge import get_active_parser_config
    parser_cfg = get_active_parser_config(db)
    llama_key = parser_cfg.get("api_key") or get_setting(db, "llama_cloud_api_key", "") or os.getenv("LLAMA_CLOUD_API_KEY", "")
    content, meta = doc.markdown_text or "", doc.cmetadata if isinstance(doc.cmetadata, dict) else {}
    file_b64, ext = meta.get("file_b64"), Path(meta.get("filename", doc.title)).suffix.lower() or ".pdf"
    if file_b64 and (not content or len(content.strip()) < 10 or meta.get("convert_to_md", True)):
        try: raw_bytes = base64.b64decode(file_b64)
        except Exception: raw_bytes = None
        if raw_bytes:
            if llama_key and meta.get("convert_to_md", True) and ext not in [".md", ".markdown", ".txt"]:
                temp_path = None
                try:
                    with tempfile.NamedTemporaryFile(delete=False, suffix=ext) as tf:
                        tf.write(raw_bytes); temp_path = tf.name
                    from llama_cloud import LlamaCloud
                    with open(temp_path, "rb") as f: file_obj = LlamaCloud(api_key=llama_key).files.create(file=f, purpose="parse")
                    args = {
                        "file_id": file_obj.id, "tier": parser_cfg.get("tier") or "agentic", "version": parser_cfg.get("version") or "latest",
                        "disable_cache": parser_cfg.get("disable_cache", False), "expand": ["markdown"],
                        "output_options": {"markdown": {"tables": {"output_tables_as_markdown": parser_cfg.get("output_tables_as_markdown", True), "compact_markdown_tables": parser_cfg.get("compact_markdown_tables", True)}}},
                    }
                    if parser_cfg.get("page_ranges"): args["page_ranges"] = parser_cfg.get("page_ranges")
                    res = LlamaCloud(api_key=llama_key).parsing.parse(**args)
                    if res and res.markdown and res.markdown.pages:
                        pages = [p.markdown for p in res.markdown.pages if p and p.success and p.markdown]
                        if pages: content = "\n\n".join(pages)
                except Exception as e: print(f"[WARN] LlamaCloud parsing error: {e}", flush=True)
                finally:
                    if temp_path and os.path.exists(temp_path):
                        try: os.remove(temp_path)
                        except Exception: pass
            if not content or len(content.strip()) < 10:
                if ext == ".pdf":
                    try:
                        import io, pypdf
                        pages = [f"## Page {i + 1}\n\n{p.extract_text().strip()}" for i, p in enumerate(pypdf.PdfReader(io.BytesIO(raw_bytes)).pages) if p.extract_text() and p.extract_text().strip()]
                        if pages: content = "\n\n".join(pages)
                    except Exception: pass
                elif ext in [".docx", ".doc"]:
                    try:
                        import docx, io
                        paras = [p.text.strip() for p in docx.Document(io.BytesIO(raw_bytes)).paragraphs if p.text.strip()]
                        if paras: content = "\n\n".join(paras)
                    except Exception: pass
                else:
                    try: content = raw_bytes.decode("utf-8")
                    except Exception: content = raw_bytes.decode("latin-1", errors="ignore")
    final_text = content or doc.content_text or doc.markdown_text or ""
    if not final_text.strip(): raise ValueError(f"Could not parse text or extract content from file '{doc.title}'. The file may be empty or corrupted.")
    return final_text

def _save_parser_output_if_configured(doc: KnowledgeDocument, content: str) -> None:
    meta = doc.cmetadata if isinstance(doc.cmetadata, dict) else {}
    if meta.get("save_parser_output", False) and content:
        try:
            parsed_dir = Path(__file__).resolve().parent.parent / "parsed_outputs"
            parsed_dir.mkdir(parents=True, exist_ok=True)
            safe_title = re.sub(r'[/\\?%*:|"<>]', "_", doc.title)
            if not safe_title.lower().endswith(".md"): safe_title += ".md"
            (parsed_dir / safe_title).write_text(content, encoding="utf-8")
        except Exception: pass

@router.post("/documents/{document_id}/ingest", response_model=DocumentResponse)
def ingest_document(document_id: uuid.UUID, current_user: User = Depends(require_admin), db: Session = Depends(get_db)):
    doc = db.query(KnowledgeDocument).filter(KnowledgeDocument.id == document_id).first()
    if not doc: raise HTTPException(status_code=404, detail="Document not found")
    col = db.query(KnowledgeCollection).filter(KnowledgeCollection.id == doc.collection_id).first()
    if not col: raise HTTPException(status_code=404, detail="Parent collection not found")
    try:
        doc.ingestion_status = "PARSING"; db.commit()
        content = _parse_content_to_markdown(doc, db)
        doc.markdown_text, doc.content_text = content, content; db.commit()
        _save_parser_output_if_configured(doc, content)
        doc.ingestion_status = "CHUNKING"; db.commit()
        doc_meta, col_meta = doc.cmetadata if isinstance(doc.cmetadata, dict) else {}, col.cmetadata if isinstance(col.cmetadata, dict) else {}
        chunk_size = doc_meta.get("chunk_size") or col_meta.get("chunk_size") or int(get_setting(db, "chunk_size", "1024"))
        chunk_overlap = doc_meta.get("chunk_overlap") if doc_meta.get("chunk_overlap") is not None else col_meta.get("chunk_overlap", int(get_setting(db, "chunk_overlap", "50")))
        from langchain_text_splitters import Language, RecursiveCharacterTextSplitter
        splitter = RecursiveCharacterTextSplitter.from_language(language=Language.MARKDOWN, chunk_size=chunk_size, chunk_overlap=chunk_overlap, length_function=len) if (doc_meta.get("convert_to_md", True) or doc.markdown_text) else RecursiveCharacterTextSplitter(separators=["\n\n", "\n", ".", "?", "!", " ", ""], chunk_size=chunk_size, chunk_overlap=chunk_overlap, length_function=len)
        chunks = splitter.split_text(content)
        doc.ingestion_status = "EMBEDDING"; db.commit()
        from config import DATABASE_URL
        from knowledge import get_embeddings
        from langchain_core.documents import Document as LCDocument
        from langchain_postgres.vectorstores import PGVector
        vector_store = PGVector(embeddings=get_embeddings(db), collection_name=col.name, connection=DATABASE_URL, use_jsonb=True, pre_delete_collection=False)
        lc_docs = [LCDocument(page_content=chunk, metadata={"document_id": str(doc.id), "collection_id": str(col.id), "title": doc.title, "document_type": doc.document_type, "chunk_index": idx}) for idx, chunk in enumerate(chunks)]
        def _purge_and_cancel():
            try:
                db.execute(text("DELETE FROM langchain_pg_embedding WHERE cmetadata->>'document_id' = :doc_id"), {"doc_id": str(document_id)})
            except Exception:
                db.rollback()
            c = db.query(KnowledgeDocument).filter(KnowledgeDocument.id == document_id).first()
            if c: db.delete(c)
            db.commit()
            now_iso = datetime.datetime.now(datetime.timezone.utc).isoformat()
            return DocumentResponse(id=str(document_id), collection_id=str(col.id), title=doc.title if doc else "Cancelled Document", document_type=doc.document_type if doc else "md", source_type="file", file_size=0, status="DRAFT", ingestion_status="CANCELLED", chunk_count=0, version=1, created_at=now_iso, updated_at=now_iso, convert_to_md=False)
        chk = db.query(KnowledgeDocument).filter(KnowledgeDocument.id == document_id).first()
        if not chk or chk.ingestion_status == "CANCELLED": return _purge_and_cancel()
        if lc_docs: vector_store.add_documents(lc_docs)
        chk = db.query(KnowledgeDocument).filter(KnowledgeDocument.id == document_id).first()
        if not chk or chk.ingestion_status == "CANCELLED": return _purge_and_cancel()
        chk.ingestion_status, chk.status, chk.chunk_count = "COMPLETED", "ACTIVE", len(chunks)
        db.commit(); db.refresh(chk); touch_collection(db, col.id)
        return make_doc_response(chk)
    except Exception as exc:
        db.rollback()
        try:
            db.execute(text("DELETE FROM langchain_pg_embedding WHERE cmetadata->>'document_id' = :doc_id"), {"doc_id": str(document_id)})
            doc_err = db.query(KnowledgeDocument).filter(KnowledgeDocument.id == document_id).first()
            if doc_err:
                if doc_err.ingestion_status == "CANCELLED": db.delete(doc_err)
                else: doc_err.ingestion_status = "FAILED"
            db.commit()
        except Exception:
            db.rollback()
        raise HTTPException(status_code=500, detail=f"Ingestion process interrupted: {exc}")

@router.post("/documents/{document_id}/cancel")
def cancel_document_ingestion(document_id: uuid.UUID, current_user: User = Depends(require_admin), db: Session = Depends(get_db)):
    doc = db.query(KnowledgeDocument).filter(KnowledgeDocument.id == document_id).first()
    if doc:
        doc.ingestion_status = "CANCELLED"; col_id = doc.collection_id
        try:
            db.execute(text("DELETE FROM langchain_pg_embedding WHERE cmetadata->>'document_id' = :doc_id"), {"doc_id": str(doc.id)})
        except Exception:
            db.rollback()
        db.delete(doc); touch_collection(db, col_id); db.commit()
    return None

@router.get("/documents/{document_id}/download")
def download_document(document_id: uuid.UUID, export_format: Optional[str] = None, current_user: User = Depends(require_admin), db: Session = Depends(get_db)):
    doc = db.query(KnowledgeDocument).filter(KnowledgeDocument.id == document_id).first()
    if not doc: raise HTTPException(status_code=404, detail="Document not found")
    meta = doc.cmetadata if isinstance(doc.cmetadata, dict) else {}
    file_b64, orig = meta.get("file_b64"), meta.get("filename") or doc.title or "document"
    if export_format == "original" and file_b64:
        try:
            content_bytes = base64.b64decode(file_b64); filename = orig
            media_type = "application/pdf" if filename.lower().endswith(".pdf") else ("application/vnd.openxmlformats-officedocument.wordprocessingml.document" if filename.lower().endswith(".docx") else ("text/plain; charset=utf-8" if filename.lower().endswith(".txt") else "application/octet-stream"))
        except Exception:
            content_bytes, filename, media_type = (doc.markdown_text or doc.content_text or "").encode("utf-8"), doc.title or "document.md", "text/markdown; charset=utf-8"
    else:
        content_bytes = (doc.markdown_text or doc.content_text or "").encode("utf-8")
        filename = doc.title or "document.md"
        if not filename.lower().endswith((".md", ".txt", ".markdown")): filename = (filename.rsplit(".", 1)[0] if "." in filename else filename) + ".md"
        media_type = "text/markdown; charset=utf-8"
    safe_ascii = re.sub(r"[^a-zA-Z0-9\.\-\_\(\) ]", "_", filename)
    encoded = urllib.parse.quote(filename, safe="")
    return Response(content=content_bytes, media_type=media_type, headers={"Content-Disposition": f'attachment; filename="{safe_ascii}"; filename*=UTF-8\'\'{encoded}', "Access-Control-Expose-Headers": "Content-Disposition"})

def _cleanup_physical_doc_files(doc: KnowledgeDocument) -> None:
    base = Path(__file__).resolve().parent.parent
    parsed_dir, uploads_dir = base / "parsed_outputs", base / "uploads"
    meta = doc.cmetadata if isinstance(doc.cmetadata, dict) else {}
    orig = meta.get("filename") or doc.title
    safe_title = re.sub(r'[/\\?%*:|"<>]', "_", doc.title)
    variants = [safe_title, f"{safe_title}.md" if not safe_title.lower().endswith(".md") else safe_title, doc.title, f"{doc.title}.md" if not doc.title.lower().endswith(".md") else doc.title, f"{doc.id}_{safe_title}"]
    if meta.get("filename"): variants.append(meta["filename"])
    if parsed_dir.exists():
        for fname in variants:
            t = parsed_dir / fname
            if t.exists() and t.is_file():
                try: t.unlink()
                except Exception: pass
    if uploads_dir.exists():
        for t in [uploads_dir / orig, uploads_dir / f"{doc.id}_{orig}"]:
            if t.exists() and t.is_file():
                try: t.unlink()
                except Exception: pass

@router.delete("/documents/{document_id}", status_code=204)
def delete_document(document_id: uuid.UUID, current_user: User = Depends(require_admin), db: Session = Depends(get_db)):
    doc = db.query(KnowledgeDocument).filter(KnowledgeDocument.id == document_id).first()
    if not doc: raise HTTPException(status_code=404, detail="Document not found")
    col_id = doc.collection_id
    _cleanup_physical_doc_files(doc)
    try: db.execute(text("DELETE FROM langchain_pg_embedding WHERE cmetadata->>'document_id' = :doc_id OR cmetadata->>'source' = :title"), {"doc_id": str(doc.id), "title": doc.title})
    except Exception: db.rollback()
    db.delete(doc); db.commit(); touch_collection(db, col_id)
    return None
